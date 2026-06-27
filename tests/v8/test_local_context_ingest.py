"""FR-013a P2 — document ingestion: extract text per format, structure into
candidate items (heuristic + coercion), and the upload API (added INACTIVE)."""
from __future__ import annotations

import io

import pytest

from portal import local_context as lc
from portal import local_context_ingest as ing


@pytest.fixture(autouse=True)
def _isolated_library(tmp_path, monkeypatch):
    monkeypatch.setattr(lc, "LIBRARY_DIR", tmp_path)
    monkeypatch.setattr(lc, "LIBRARY_PATH", tmp_path / "library.json")


# ── extract_text per format ──────────────────────────────────────────────────

def test_extract_txt():
    assert ing.extract_text("notes.txt", b"line one\nline two") == "line one\nline two"


def test_extract_xlsx():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Drug", "Rule"]); ws.append(["Norepinephrine", "max 30 mcg/min"])
    buf = io.BytesIO(); wb.save(buf)
    text = ing.extract_text("formulary.xlsx", buf.getvalue())
    assert "Norepinephrine" in text and "max 30 mcg/min" in text


def test_extract_docx():
    import docx
    d = docx.Document(); d.add_paragraph("Lactate q2h until normalized.")
    buf = io.BytesIO(); d.save(buf)
    text = ing.extract_text("orders.docx", buf.getvalue())
    assert "Lactate q2h until normalized." in text


def test_extract_unsupported_raises():
    with pytest.raises(ValueError):
        ing.extract_text("archive.zip", b"PK\x03\x04")


# ── parse_items (heuristic + coercion) ───────────────────────────────────────

def test_parse_heuristic_line_split_no_key():
    items = ing.parse_items("- Lactate q2h until normalized.\n* Norepi first-line\nx", api_key="")
    assert len(items) == 2                                   # 'x' too short, dropped
    assert all(i["type"] == "standing_order" for i in items)
    assert items[0]["content"] == "Lactate q2h until normalized."
    assert items[0]["title"]                                 # derived


def test_parse_coerces_ai_items(monkeypatch):
    monkeypatch.setattr(ing, "_ai_items", lambda text, key: [
        {"type": "bogus", "title": "T", "content": "C1"},        # bad type → standing_order
        {"type": "medication", "title": "", "content": "Pressor of choice"},  # title derived
        {"type": "standing_order", "title": "x", "content": ""}, # empty content → dropped
        "not-a-dict",                                            # skipped
    ])
    items = ing.parse_items("whatever", api_key="sk-dummy")
    assert [i["type"] for i in items] == ["standing_order", "medication"]
    assert items[1]["title"] == "Pressor of choice"             # derived from content


# ── ingest() adds candidates INACTIVE ────────────────────────────────────────

def test_ingest_adds_inactive_candidates():
    res = ing.ingest("orders.txt",
                     b"- Lactate q2h until normalized.\n- Norepi first-line pressor", api_key="")
    assert res["added"] == 2
    items = lc.list_items()
    assert len(items) == 2
    assert all(it["active"] is False for it in items)          # nothing live until confirmed
    assert all(it["source"] == "orders.txt" for it in items)


# ── ingest API ───────────────────────────────────────────────────────────────

@pytest.fixture
def client(tmp_path, monkeypatch):
    fake_home = tmp_path / "home"
    fake_home.mkdir()
    monkeypatch.setenv("HOME", str(fake_home))
    monkeypatch.setenv("MEDSIM_RESUME", "0")
    from portal import auth, control_room, credentials, server as server_mod
    sb = fake_home / ".medsim"
    sb.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(credentials, "VAULT_DIR", sb)
    monkeypatch.setattr(credentials, "VAULT_PATH", sb / "vault.enc")
    monkeypatch.setattr(server_mod, "_anthropic_runtime_key", "")
    control_room._reset_for_tests()
    if not credentials.is_initialized():
        credentials.initialize("test_passwd_xyz_8chars")
    vault = credentials.unlock("test_passwd_xyz_8chars")
    vault.set("ANTHROPIC_API_KEY", "sk-ant-dummy")
    from portal import server
    from fastapi.testclient import TestClient
    with TestClient(server.app) as c:
        c.cookies.set(auth.COOKIE_NAME, auth.issue_session_token(vault))
        yield c
    control_room._reset_for_tests()


def test_api_ingest_txt_adds_inactive(client, monkeypatch):
    # force the heuristic path (no live Anthropic call with the dummy key)
    monkeypatch.setattr(ing, "_ai_items",
                        lambda t, k: (_ for _ in ()).throw(RuntimeError("offline")))
    files = {"file": ("orders.txt", b"- Lactate q2h\n- Norepi first-line", "text/plain")}
    r = client.post("/api/local-context/ingest", files=files)
    assert r.status_code == 200, r.text
    assert r.json()["added"] == 2
    body = client.get("/api/local-context/items").json()
    assert len(body["items"]) == 2
    assert all(not it["active"] for it in body["items"])       # inactive, awaiting review


def test_api_ingest_unsupported_type_400(client):
    files = {"file": ("x.zip", b"PK\x03\x04", "application/zip")}
    assert client.post("/api/local-context/ingest", files=files).status_code == 400


def test_api_ingest_requires_instructor_auth(client):
    from portal import auth
    client.cookies.delete(auth.COOKIE_NAME)
    files = {"file": ("x.txt", b"a substantive line here", "text/plain")}
    assert client.post("/api/local-context/ingest", files=files).status_code in (401, 403)
